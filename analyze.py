#!/usr/bin/env python3
"""Разбор любого шаблона .pptx: какие в нём макеты (layouts), какие поля и сколько текста в них помещается.

    python3 analyze.py template.pptx -o my-template/layouts.json            # каталог макетов
    python3 analyze.py template.pptx -o my-template/layouts.json --render   # + карта макетов (картинки)

Результат — layouts.json в формате, который понимает build.py. Его стоит просмотреть: названия полей и описания
макетов берутся из шаблона и эвристик, их можно уточнить вручную.
"""
import argparse
import copy
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.oxml.ns import qn
except ImportError:
    sys.exit('Нужен python-pptx: pip install python-pptx')

META = ('DATE', 'FOOTER', 'SLIDE_NUMBER')
SKIP_LAYOUT = re.compile(r'vertical|вертикальн', re.I)


def ph_type(ph):
    return str(ph.placeholder_format.type).split(' ')[0].split('.')[-1]


def _xfrm(el):
    sp = el.find(qn('p:spPr'))
    x = sp.find(qn('a:xfrm')) if sp is not None else None
    if x is None or x.find(qn('a:off')) is None:
        return None
    o, e = x.find(qn('a:off')), x.find(qn('a:ext'))
    return int(o.get('x')), int(o.get('y')), int(e.get('cx')), int(e.get('cy'))


def _size(txBody):
    for el in txBody.iter():
        if el.tag in (qn('a:rPr'), qn('a:endParaRPr'), qn('a:defRPr')) and el.get('sz'):
            return int(el.get('sz')) / 100
    return None


def _same_kind(a, b):
    t = lambda x: 'title' if x in ('TITLE', 'CENTER_TITLE') else 'body' if x in ('BODY', 'OBJECT') else x
    return t(a) == t(b)


def inherited(layout, ph):
    """bbox (EMU) и кегль (pt) плейсхолдера макета с учётом мастера."""
    bbox = _xfrm(ph._element)
    tb = ph._element.find(qn('p:txBody'))
    fs = _size(tb) if tb is not None else None
    master = layout.slide_master
    if bbox is None or fs is None:
        for mph in master.placeholders:
            if _same_kind(ph_type(mph), ph_type(ph)) or mph.placeholder_format.idx == ph.placeholder_format.idx:
                if bbox is None:
                    bbox = _xfrm(mph._element)
                mtb = mph._element.find(qn('p:txBody'))
                if fs is None and mtb is not None:
                    fs = _size(mtb)
                break
    if fs is None:
        st = master._element.find(qn('p:txStyles'))
        if st is not None:
            node = st.find(qn('p:titleStyle') if ph_type(ph) in ('TITLE', 'CENTER_TITLE') else qn('p:bodyStyle'))
            d = node.find('.//' + qn('a:lvl1pPr') + '/' + qn('a:defRPr')) if node is not None else None
            if d is not None and d.get('sz'):
                fs = int(d.get('sz')) / 100
    return bbox, fs or 18


def has_bullets(layout, ph):
    """Есть ли у поля маркеры списка: стиль плейсхолдера макета -> мастера -> bodyStyle мастера."""
    def bullet_of(txBody):
        if txBody is None:
            return None
        lvl = txBody.find(qn('a:lstStyle') + '/' + qn('a:lvl1pPr'))
        if lvl is None:
            p = txBody.find(qn('a:p') + '/' + qn('a:pPr'))
            lvl = p
        if lvl is None:
            return None
        if lvl.find(qn('a:buNone')) is not None:
            return False
        if lvl.find(qn('a:buChar')) is not None or lvl.find(qn('a:buAutoNum')) is not None:
            return True
        return None
    b = bullet_of(ph._element.find(qn('p:txBody')))
    if b is not None:
        return b
    master = layout.slide_master
    for mph in master.placeholders:
        if _same_kind(ph_type(mph), ph_type(ph)):
            b = bullet_of(mph._element.find(qn('p:txBody')))
            if b is not None:
                return b
            break
    st = master._element.find(qn('p:txStyles'))
    node = st.find(qn('p:bodyStyle') + '/' + qn('a:lvl1pPr')) if st is not None else None
    if node is not None:
        if node.find(qn('a:buNone')) is not None:
            return False
        return node.find(qn('a:buChar')) is not None or node.find(qn('a:buAutoNum')) is not None
    return True


def capacity(bbox, fs):
    w = max(bbox[2] / 12700 - 14, fs)
    h = max(bbox[3] / 12700 - 7, fs * 1.2)
    per_line = max(1, int(w / (fs * 0.55)))
    lines = max(1, int(h / (fs * 1.2)))
    return per_line, lines


GENERIC_NAMES = re.compile(r'^(placeholder|заполнитель|место для|text placeholder|content placeholder|'
                           r'title|заголовок|subtitle|подзаголовок|body|текст|объект|object)\b', re.I)
TYPE_LABELS = {'TITLE': 'Заголовок', 'CENTER_TITLE': 'Заголовок', 'SUBTITLE': 'Подзаголовок', 'BODY': 'Текст',
               'OBJECT': 'Текст', 'PICTURE': 'Рисунок'}


def label(ph):
    name = re.sub(r'\s*\d+$', '', ph.name).strip()
    if not name or re.match(r'^(placeholder|заполнитель)', name, re.I):
        return TYPE_LABELS.get(ph_type(ph), ph_type(ph).lower())
    return name


def safe(n):
    return max(1, int(n * 0.75))


def analyze(path):
    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    catalog = {}
    seen = {}
    for layout in prs.slide_layouts:
        if SKIP_LAYOUT.search(layout.name):
            continue
        phs = [ph for ph in layout.placeholders if ph_type(ph) not in META]
        if not phs:
            continue
        info = []
        for ph in phs:
            bbox, fs = inherited(layout, ph)
            if bbox is None:
                continue
            cx, cy = (bbox[0] + bbox[2] / 2) / W, (bbox[1] + bbox[3] / 2) / H
            per_line, lines = capacity(bbox, fs)
            info.append(dict(ph=ph, idx=ph.placeholder_format.idx, type=ph_type(ph), bbox=bbox, fs=fs,
                             cx=cx, cy=cy, w=bbox[2] / W, h=bbox[3] / H, per_line=per_line, lines=lines))
        fields, items = {}, None
        titles = [i for i in info if i['type'] in ('TITLE', 'CENTER_TITLE')]
        if titles:
            t = titles[0]
            fields['title'] = {'idx': t['idx'], 'ph_type': 'TITLE', 'type': 'text', 'label': label(t['ph']) or 'Заголовок',
                               'max': safe(t['per_line'] * min(t['lines'], 3)), 'required': True}
        for i in info:
            if i['type'] == 'SUBTITLE':
                fields['subtitle'] = {'idx': i['idx'], 'ph_type': 'SUBTITLE', 'type': 'text', 'label': label(i['ph']),
                                      'max': safe(i['per_line'] * i['lines'])}
            if i['type'] == 'PICTURE':
                fields['image' if 'image' not in fields else f'image{i["idx"]}'] = {
                    'idx': i['idx'], 'ph_type': 'PICTURE', 'type': 'image', 'label': label(i['ph']) + ' (путь к файлу картинки)'}
        bodies = [i for i in info if i['type'] in ('BODY', 'OBJECT')]
        # группы: несколько одинаковых по ширине полей в ряд (карточки, колонки, плитки цифр)
        cols = []
        for b in sorted(bodies, key=lambda b: b['cx']):
            for c in cols:
                if abs(c[0]['cx'] - b['cx']) < 0.05 and abs(c[0]['w'] - b['w']) < 0.05:
                    c.append(b)
                    break
            else:
                cols.append([b])
        group = [c for c in cols if len(c) == len(cols[0])] if cols else []
        if len(group) >= 2 and len(group) == len(cols) and all(abs(c[0]['w'] - group[0][0]['w']) < 0.05 for c in group):
            rows = len(group[0])
            for c in group:
                c.sort(key=lambda b: b['cy'])
            names = ['title', 'text'] if rows == 2 else (['text'] if rows == 1 else [f'f{r + 1}' for r in range(rows)])
            ifields = {}
            for r in range(rows):
                cells = [c[r] for c in group]
                lab = label(cells[0]['ph'])
                is_list = cells[0]['lines'] >= 3 and has_bullets(layout, cells[0]['ph'])
                f = {'idx': [x['idx'] for x in cells], 'ph_type': [x['type'] for x in cells], 'type': 'text', 'label': lab,
                     'max': safe(min(x['per_line'] * x['lines'] for x in cells)), 'required': True}
                if is_list:
                    f.update(type='list', max_items=min(5, cells[0]['lines']), max=safe(cells[0]['per_line'] * 2),
                             max_total=safe(cells[0]['per_line'] * cells[0]['lines']))
                ifields[names[r]] = f
            # декоративные фигуры макета под колонками (карточки) — тогда нужно ровно n элементов
            decor = 0
            for el in layout.shapes:
                if el.is_placeholder or el.width is None:
                    continue
                ecx = (el.left + el.width / 2) / W
                if any(abs(ecx - c[0]['cx']) < 0.05 for c in group) and el.width / W < 0.6:
                    decor += 1
            n = len(group)
            items = {'key': 'items', 'label': 'Элементы (колонки/карточки/плитки)',
                     'min': n if decor >= n else min(2, n), 'max': n, 'fields': ifields}
        else:
            for k, b in enumerate(sorted(bodies, key=lambda b: (round(b['cy'], 1), b['cx']))):
                key = 'body' if k == 0 else f'body{k + 1}'
                if b['lines'] >= 3 and has_bullets(layout, b['ph']):  # маркеры в стиле поля — значит, список
                    fields[key] = {'idx': b['idx'], 'ph_type': b['type'], 'type': 'list', 'label': label(b['ph']),
                                   'max_items': min(6, b['lines']), 'max': safe(b['per_line'] * 2),
                                   'max_total': safe(b['per_line'] * b['lines']), 'required': k == 0}
                else:
                    fields[key] = {'idx': b['idx'], 'ph_type': b['type'], 'type': 'text', 'label': label(b['ph']),
                                   'max': safe(b['per_line'] * b['lines'])}
        has_table = any(sh.has_table for s in prs.slides if s.slide_layout is layout for sh in s.shapes)
        if has_table or (set(fields) == {'title'} and not items):
            fields['table'] = {'type': 'table', 'label': 'Таблица: header — заголовки столбцов, rows — строки',
                               'max_cols': 6, 'max_rows': 7, 'max': 35}
        sig = json.dumps([fields, items], sort_keys=True, default=str)
        if sig in seen:  # дубликаты макетов (бывают в шаблонах) пропускаем
            continue
        seen[sig] = layout.name
        desc = []
        if 'title' in fields:
            desc.append('заголовок')
        if 'subtitle' in fields:
            desc.append('подзаголовок')
        if items:
            desc.append(f'{items["max"]} элемента в ряд ({", ".join(items["fields"])})')
        for k, f in fields.items():
            if k.startswith('body'):
                desc.append('список' if f['type'] == 'list' else 'текст')
            if f['type'] == 'image':
                desc.append('картинка')
            if f['type'] == 'table':
                desc.append('таблица')
        entry = {'use': ' + '.join(desc), 'fields': fields}
        if items:
            entry['items'] = items
        catalog[layout.name] = entry
    return {
        '_about': f'Каталог макетов шаблона {Path(path).name}, собран analyze.py. Проверь названия и описания (use).',
        'template': Path(path).name,
        'layouts': catalog,
        'rules': [
            'Первый слайд — титульный макет шаблона, последний — финальный (если такой есть).',
            'Для группы элементов (карточки, колонки) передавай столько элементов, сколько допускает макет (min–max).',
            'Не больше трёх слайдов подряд на одном макете: чередуй подачу.',
            'Тексты — только из исходного материала, без новых фактов и цифр.',
            'Укладывайся в лимиты знаков: длинное — сократи или разбей на два слайда.'],
    }


def design_in_slides(prs):
    """Фирменный дизайн нарисован прямо на слайдах, а макеты пустые? -> (да/нет, декор на слайд, декор в макетах)."""
    def decor(shapes):
        return sum(1 for sh in shapes if not sh.is_placeholder)
    per_slide = (sum(decor(s.shapes) for s in prs.slides) / len(prs.slides)) if len(prs.slides) else 0
    in_layouts = sum(decor(l.shapes) for l in prs.slide_layouts) + sum(decor(m.shapes) for m in prs.slide_masters)
    per_layout = in_layouts / max(1, len(prs.slide_layouts))
    return per_slide >= 3 and per_slide > 3 * per_layout, round(per_slide, 1), round(per_layout, 1)


def render_map(template, catalog, outdir):
    """Слайд на каждый макет, поля подписаны ключами — чтобы агент видел, где какое поле."""
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice or not shutil.which('pdftoppm'):
        print('Карта макетов пропущена: нет LibreOffice или pdftoppm')
        return
    prs = Presentation(str(template))
    originals = list(prs.slides)
    by_name = {l.name: l for l in prs.slide_layouts}
    for name, cfg in catalog['layouts'].items():
        s = prs.slides.add_slide(by_name[name])
        pl = list(s.placeholders)

        def pick(idx, t):
            c = [p for p in pl if p.placeholder_format.idx == idx]
            if len(c) > 1 and t:
                c = [p for p in c if ph_type(p).replace('CENTER_', '') == t.replace('CENTER_', '')] or c
            return c[0] if c else None
        for k, f in cfg['fields'].items():
            if 'idx' in f:
                p = pick(f['idx'], f.get('ph_type'))
                if p is not None and p.has_text_frame:
                    p.text_frame.text = f'[{k}] «{name}»' if k == 'title' else f'[{k}]'
        if 'items' in cfg:
            for fk, f in cfg['items']['fields'].items():
                for i, idx in enumerate(f['idx']):
                    p = pick(idx, f['ph_type'][i] if isinstance(f.get('ph_type'), list) else None)
                    if p is not None and p.has_text_frame:
                        p.text_frame.text = f'[items[{i}].{fk}]'
    lst = prs.slides._sldIdLst
    for sl in originals:
        for sid in list(lst):
            if prs.part.related_part(sid.rId) is sl.part:
                rid = sid.rId
                lst.remove(sid)
                prs.part.drop_rel(rid)
    outdir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        prs.save(f'{td}/map.pptx')
        subprocess.run([soffice, f'-env:UserInstallation=file://{td}/lo', '--headless', '--convert-to', 'pdf',
                        '--outdir', td, f'{td}/map.pptx'], capture_output=True, timeout=600)
        if not Path(f'{td}/map.pdf').exists():
            print('Карта макетов не отрендерилась')
            return
        for f in outdir.glob('layout-*.png'):
            f.unlink()
        subprocess.run(['pdftoppm', '-r', '50', '-png', f'{td}/map.pdf', str(outdir / 'layout')], timeout=600)
    pngs = sorted(outdir.glob('layout-*.png'), key=lambda p: int(re.findall(r'(\d+)\.png$', p.name)[0]))
    try:
        from PIL import Image, ImageDraw
        ims = [Image.open(p).convert('RGB') for p in pngs]
        w = 420
        ims = [im.resize((w, int(im.height * w / im.width))) for im in ims]
        cols, h = 3, ims[0].height
        names = list(catalog['layouts'])
        sheet = Image.new('RGB', (cols * (w + 10) + 10, ((len(ims) + cols - 1) // cols) * (h + 26) + 10), 'white')
        d = ImageDraw.Draw(sheet)
        for i, im in enumerate(ims):
            x, y = 10 + (i % cols) * (w + 10), 10 + (i // cols) * (h + 26)
            d.text((x, y), f'{i + 1}. {names[i] if i < len(names) else ""}', fill='black')
            sheet.paste(im, (x, y + 14))
        sheet.save(outdir / 'layouts-map.jpg', quality=85)
        print(f'Карта макетов: {outdir / "layouts-map.jpg"}')
    except Exception as e:
        print(f'Карта макетов: картинки в {outdir} (обзор не собран: {e})')


def main():
    ap = argparse.ArgumentParser(description='Каталог макетов любого шаблона .pptx для build.py')
    ap.add_argument('template')
    ap.add_argument('-o', '--out', required=True, help='куда записать layouts.json')
    ap.add_argument('--render', action='store_true', help='нарисовать карту макетов с подписями полей')
    a = ap.parse_args()
    cat = analyze(a.template)
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(cat, ensure_ascii=False, indent=1), encoding='utf-8')
    flag, per_slide, per_layout = design_in_slides(Presentation(a.template))
    print(f'Каталог: {out} — макетов: {len(cat["layouts"])}')
    for name, cfg in cat['layouts'].items():
        print(f'  «{name}»: {cfg["use"]}')
    if flag:
        print(f'ВНИМАНИЕ: фирменный дизайн нарисован прямо на слайдах шаблона (в среднем {per_slide} декоративных '
              f'объектов на слайд против {per_layout} в макете). Сборка по макетам даст безликие слайды. Используй '
              'клонирование слайдов шаблона: https://github.com/sergeykruglovit-a11y/owui-pptx-kit')
    if not cat['layouts']:
        print('ВНИМАНИЕ: в шаблоне нет пригодных макетов с полями — используй https://github.com/sergeykruglovit-a11y/owui-pptx-kit')
    if a.render:
        render_map(a.template, cat, out.parent / 'layouts-map')


if __name__ == '__main__':
    main()
